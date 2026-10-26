"""
tests/test_resume_tailor.py — smoke tests for resume_tailor.py.

Run:  python3 -m tests.test_resume_tailor
"""

import os
import re
import sys
import shutil
import tempfile
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))


# ── Self-contained: set up a fake source resume BEFORE importing
#    resume_tailor, because resume_tailor.py reads RESUME_DIR at import
#    time (it's a module-level constant, not a per-call lookup).
#    On the GitHub Actions Ubuntu runner the default RESUME_DIR
#    (~/PycharmProjects/data-etl-ml/Resume) doesn't exist, which made
#    tailor_resume() return None and the test exit 1 — failing the
#    leads-ci job for the last 5 runs.
_RESUME_FAKE_DIR = Path(tempfile.mkdtemp(prefix="test_resume_"))
os.environ["RESUME_DIR"] = str(_RESUME_FAKE_DIR)
# Create a minimal docx that python-docx can open + write back.
try:
    from docx import Document
    _doc = Document()
    _doc.add_paragraph("Test headline paragraph")
    _doc.add_paragraph("Test summary paragraph two for length padding.")
    _doc.save(str(_RESUME_FAKE_DIR / "Pv_Cv_Data_Ai_2026.docx"))
except ImportError:
    pass  # docx not available — test_tailor_resume_writes_file will skip

from resume_tailor import (  # noqa: E402
    classify_jd, HEADLINES, SUMMARIES, SKILL_RANK_HINTS,
    tailor_resume, HAVE_DOCX,
)


CASES = [
    # (label, jd_text, role, expected_primary)
    ("RAG heavy",        "Production RAG, vector embeddings, GPT-4",       "Senior AI Eng",    "llm_ai"),
    ("Agentic AI",       "agentic workflow, multi-agent, tool use",       "AI Engineer",      "agentic_ai"),
    ("MLOps / serving",  "MLOps platform, model serving, KServe",         "ML Platform Eng",  "ml_platform"),
    ("Data lakehouse",   "Databricks lakehouse, Spark, Kafka, CDC",       "Data Engineer",    "data_platform"),
    ("Cloud arch",       "AWS, Azure, Kubernetes, multi-cloud, Terraform","Cloud Architect", "cloud_arch"),
    ("GenAI app",        "GenAI chatbot, RAG, summarization",             "GenAI Eng",        "genai_app"),
    ("Full-stack",       "React, Node.js, TypeScript, full-stack",        "Full-Stack Eng",   "full_stack"),
    ("Leadership",       "Head of Data, manage team, $2M budget",         "Head of Data",     "leadership"),
    ("Empty / generic",  "Engineer wanted",                              "Engineer",         "data_platform"),  # default
]


def test_classify():
    print("🧪 test_classify_jd")
    fails = 0
    for label, jd, role, want in CASES:
        got = classify_jd(jd, role)
        primary = got[0] if got else None
        ok = primary == want
        marker = "✅" if ok else "❌"
        if not ok:
            fails += 1
        print(f"  {marker} {label:18s} primary={primary:14s} want={want:14s} all={got}")
    print(f"  → {len(CASES) - fails}/{len(CASES)} pass")
    return fails == 0


def test_headline_and_summary_for_each_family():
    print("\n🧪 test_headline_and_summary_for_each_family")
    fails = 0
    for family in HEADLINES:
        if family not in SUMMARIES or family not in SKILL_RANK_HINTS:
            print(f"  ❌ {family:14s} missing SUMMARIES or SKILL_RANK_HINTS entry")
            fails += 1
            continue
        h = HEADLINES[family]
        s = SUMMARIES[family]
        if not h or len(h) < 10:
            print(f"  ❌ {family:14s} headline too short: {h!r}")
            fails += 1
        if not s or len(s) < 50:
            print(f"  ❌ {family:14s} summary too short")
            fails += 1
        if not SKILL_RANK_HINTS[family]:
            print(f"  ❌ {family:14s} empty SKILL_RANK_HINTS")
            fails += 1
        else:
            print(f"  ✅ {family:14s} headline={len(h)}c  summary={len(s)}c  skills={len(SKILL_RANK_HINTS[family])}")
    return fails == 0


def test_tailor_resume_writes_file():
    if not HAVE_DOCX:
        print("\n🧪 test_tailor_resume_writes_file: SKIPPED (no python-docx)")
        return True
    print("\n🧪 test_tailor_resume_writes_file")
    from resume_tailor import TAILORED_DIR
    pre_count = len(list(TAILORED_DIR.glob("*.docx")))
    p = tailor_resume({
        "company": "TestCo",
        "role": "Senior AI Engineer (RAG)",
        "jd_text": "Production RAG on LangChain + Pinecone, multi-agent, GPT-4",
    })
    if not p or not p.exists():
        print(f"  ❌ tailor_resume returned {p!r}")
        return False
    print(f"  ✅ wrote {p.name} ({p.stat().st_size} bytes)")
    # Verify the file is a real docx (zip starts with PK)
    with open(p, "rb") as f:
        magic = f.read(2)
    if magic != b"PK":
        print(f"  ❌ not a valid docx (magic={magic!r})")
        return False
    print(f"  ✅ file is a valid docx (PK magic)")
    # Clean up: remove the test file so it doesn't accumulate
    p.unlink(missing_ok=True)
    return True


def main():
    r1 = test_classify()
    r2 = test_headline_and_summary_for_each_family()
    r3 = test_tailor_resume_writes_file()
    all_ok = r1 and r2 and r3
    # Clean up the temp resume dir created at import time
    try:
        shutil.rmtree(_RESUME_FAKE_DIR, ignore_errors=True)
    except Exception:
        pass
    print(f"\n{'✅ ALL PASS' if all_ok else '❌ SOME FAILED'}")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
